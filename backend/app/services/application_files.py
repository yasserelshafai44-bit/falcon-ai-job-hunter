"""Downloadable DOCX preserving approved material text; no regeneration."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def resume_docx(content: str) -> bytes:
    # standard_business_brief; CV override: black headings, compact 14pt section
    # headings, plain left-aligned masthead. Text is never regenerated or amended.
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    for name, size, before, after in (
        ("Normal", 11, 0, 6),
        ("Title", 22, 0, 8),
        ("Heading 1", 14, 16, 8),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 12, 8, 4),
        ("List Bullet", 11, 0, 8),
    ):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.167 if name == "List Bullet" else 1.1
        style.paragraph_format.widow_control = True
        if style.element.pPr is not None:
            for border in list(style.element.pPr.findall(qn("w:pBdr"))):
                style.element.pPr.remove(border)
    for indent in doc.part.numbering_part.element.iter(qn("w:ind")):
        indent.set(qn("w:left"), "720")
        indent.set(qn("w:hanging"), "360")
    headings = {
        "professional summary",
        "selected operations experience and results",
        "professional experience",
        "work history",
        "education",
        "qualifications",
        "core skills",
        "key achievements",
        "selected achievements",
    }
    lines = [line.strip() for line in content.splitlines() if line.strip()]
    for index, line in enumerate(lines):
        if index == 0:
            doc.add_paragraph(line, "Title")
        elif line.casefold() in headings:
            doc.add_paragraph(line, "Heading 1")
        elif line.startswith(("- ", "• ")):
            doc.add_paragraph(line[2:], "List Bullet")
        else:
            doc.add_paragraph(line)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    buffer = BytesIO()
    doc.save(buffer)
    output = BytesIO()
    with ZipFile(buffer) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return output.getvalue()
